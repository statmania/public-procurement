#!/usr/bin/env python3
"""Convert .pptx files to Markdown, extracting pictures.

Usage: pptx_to_md.py SRC_DIR [OUT_DIR]   (default OUT_DIR: resources)

Writes OUT_DIR/<deck>.md and OUT_DIR/img/<deck>/slide<N>_<M>.<ext>.
Searches SRC_DIR recursively; skips PowerPoint lock files (~$*).
"""
import re
import sys
from pathlib import Path
from urllib.parse import quote

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BOILERPLATE = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}


def walk(shapes):
    """Yield leaf shapes in reading order, descending into groups."""
    for sh in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def is_title(sh):
    return sh.is_placeholder and sh.placeholder_format.type in TITLE_TYPES


def para_lines(tf):
    for p in tf.paragraphs:
        text = re.sub(r"[\x00-\x1f]+", " ", "".join(r.text for r in p.runs)).strip()
        if text:
            yield "  " * p.level + "- " + text


def table_md(tbl):
    rows = [[c.text.strip().replace("\n", " ").replace("|", "\\|") for c in r.cells] for r in tbl.rows]
    if not rows:
        return []
    out = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * len(rows[0])]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return out


def convert(pptx_path, out_dir):
    deck = pptx_path.stem
    img_dir = out_dir / "img" / deck
    prs = Presentation(pptx_path)
    md = [f"# {deck}", "", f"_Source: `{pptx_path.name}`_", ""]
    n_img = 0
    for n, slide in enumerate(prs.slides, 1):
        shapes = list(walk(slide.shapes))
        title = next((s.text_frame.text.strip() for s in shapes if is_title(s) and s.has_text_frame), "")
        md += [f"## Slide {n}" + (f": {title}" if title else ""), ""]
        m = 0
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                m += 1
                img_dir.mkdir(parents=True, exist_ok=True)
                name = f"slide{n}_{m}.{sh.image.ext}"
                (img_dir / name).write_bytes(sh.image.blob)
                md += [f"![](img/{quote(deck)}/{name})", ""]
                n_img += 1
            elif sh.has_table if hasattr(sh, "has_table") else False:
                md += table_md(sh.table) + [""]
            elif sh.has_text_frame and not is_title(sh) and not (sh.is_placeholder and sh.placeholder_format.type in BOILERPLATE):
                lines = list(para_lines(sh.text_frame))
                if lines:
                    md += lines + [""]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                md += [f"> **Notes:** {re.sub(chr(10) + '+', ' ', notes)}", ""]
    (out_dir / f"{deck}.md").write_text("\n".join(md), encoding="utf-8")
    return len(prs.slides), n_img


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "resources")
    out.mkdir(parents=True, exist_ok=True)
    seen = {}
    for f in sorted(src.rglob("*.pptx")):
        if f.name.startswith("~$"):
            continue
        if f.stem in seen:
            print(f"SKIP duplicate name: {f} (same as {seen[f.stem]})")
            continue
        seen[f.stem] = f
        try:
            slides, imgs = convert(f, out)
            print(f"{f.stem}: {slides} slides, {imgs} images")
        except Exception as e:
            print(f"FAIL {f}: {e}")


if __name__ == "__main__":
    main()
